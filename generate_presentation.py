"""Generate DMACO presentation (PPTX) with T-Mobile branding."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Brand colours ─────────────────────────────────────────────────
MAGENTA   = RGBColor(0xE2, 0x00, 0x74)
DARK_BG   = RGBColor(0x1A, 0x1A, 0x2E)
SURFACE   = RGBColor(0x23, 0x23, 0x3A)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY= RGBColor(0xA0, 0xA0, 0xB8)
DARK_GRAY = RGBColor(0x2D, 0x2D, 0x44)


def _set_slide_bg(slide, color=DARK_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_shape_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def _add_text_box(slide, left, top, width, height, text, font_size=18,
                  bold=False, color=WHITE, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def _add_bullet_list(slide, left, top, width, height, items, font_size=16,
                     color=WHITE, bullet_color=MAGENTA):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(8)
        p.level = 0
    return txBox


def _section_header(slide, text):
    """Magenta accent bar + section title."""
    _add_shape_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.08), MAGENTA)
    _add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.8),
                  text, font_size=32, bold=True, color=MAGENTA)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 1 — Title
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_slide_bg(slide)
    _add_shape_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), MAGENTA)

    _add_text_box(slide, Inches(1), Inches(1.8), Inches(11), Inches(1.2),
                  "DMACO", font_size=54, bold=True, color=MAGENTA)
    _add_text_box(slide, Inches(1), Inches(3.0), Inches(11), Inches(0.8),
                  "Deterministic Multi-Agent Configuration Orchestrator",
                  font_size=26, bold=False, color=WHITE)
    _add_text_box(slide, Inches(1), Inches(4.0), Inches(11), Inches(0.6),
                  "Guided AI-powered configuration for Text-to-SQL bots",
                  font_size=18, color=LIGHT_GRAY)
    _add_shape_rect(slide, Inches(1), Inches(5.0), Inches(3), Inches(0.04), MAGENTA)
    _add_text_box(slide, Inches(1), Inches(5.3), Inches(11), Inches(0.5),
                  "T-Mobile  |  April 2026", font_size=16, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 2 — Agenda
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "Agenda")
    _add_bullet_list(slide, Inches(1), Inches(1.8), Inches(10), Inches(4), [
        "1.  The Problem — Why configuring Text-to-SQL bots is painful today",
        "2.  Our Approach — DMACO architecture & hybrid AI strategy",
        "3.  Live Demo — Telecom billing bot configuration, end to end",
    ], font_size=22)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 3 — The Problem (overview)
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "1. The Problem")
    _add_text_box(slide, Inches(1), Inches(1.8), Inches(11), Inches(0.7),
                  "Configuring Text-to-SQL bots today is manual, error-prone, and slow.",
                  font_size=22, color=WHITE)
    _add_bullet_list(slide, Inches(1), Inches(2.8), Inches(5.5), Inches(4), [
        "Engineers hand-craft JSON configs with 100+ fields",
        "No validation until deploy — typos break production",
        "Domain experts can't participate (too technical)",
        "Every new domain means starting from scratch",
        "No audit trail of configuration changes",
    ], font_size=18, color=LIGHT_GRAY)

    # Right side — "pain" stats card
    card = _add_shape_rect(slide, Inches(7.5), Inches(2.6), Inches(4.8), Inches(3.5), SURFACE)
    _add_text_box(slide, Inches(7.8), Inches(2.8), Inches(4.2), Inches(0.5),
                  "Typical pain points", font_size=18, bold=True, color=MAGENTA)
    _add_bullet_list(slide, Inches(7.8), Inches(3.5), Inches(4.2), Inches(2.5), [
        "~2-4 hours per configuration",
        "~30% configs have errors on first deploy",
        "0 visibility for business stakeholders",
        "No reusable templates across domains",
    ], font_size=16, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 4 — The Problem (deeper)
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "1. The Problem — What Goes Wrong")

    cols = [
        ("Manual JSON Editing", [
            "Deeply nested structure",
            "Easy to mistype field names",
            "No schema enforcement at edit time",
        ]),
        ("No Guided Workflow", [
            "Users jump between docs and editor",
            "Steps can be skipped or done out of order",
            "Onboarding takes days, not minutes",
        ]),
        ("Zero Change Tracking", [
            "Who edited what? When?",
            "No diff between config versions",
            "Impossible to audit for compliance",
        ]),
    ]
    for idx, (title, bullets) in enumerate(cols):
        x = Inches(1 + idx * 4)
        _add_shape_rect(slide, x, Inches(1.8), Inches(3.5), Inches(4.2), SURFACE)
        _add_text_box(slide, x + Inches(0.3), Inches(1.95), Inches(3), Inches(0.5),
                      title, font_size=18, bold=True, color=MAGENTA)
        _add_bullet_list(slide, x + Inches(0.3), Inches(2.6), Inches(3), Inches(3),
                         bullets, font_size=15, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 5 — Our Approach (overview)
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "2. Our Approach — DMACO")

    _add_text_box(slide, Inches(1), Inches(1.8), Inches(11), Inches(0.7),
                  "A hybrid AI + deterministic pipeline that guides users step-by-step.",
                  font_size=22, color=WHITE)

    principles = [
        ("Finite-State Machine", "15-state FSM enforces correct order.\nNo skipping, no invalid transitions."),
        ("7 Specialized AI Agents", "Each step has its own LLM agent\nwith isolated system prompts."),
        ("Deterministic Validators", "Every LLM output is gated by\nrule-based validation before storage."),
        ("Chat-Based UI", "Non-technical users can configure\nbots through natural conversation."),
    ]
    for idx, (title, desc) in enumerate(principles):
        x = Inches(0.6 + idx * 3.15)
        _add_shape_rect(slide, x, Inches(3.0), Inches(2.9), Inches(3.2), SURFACE)
        _add_text_box(slide, x + Inches(0.2), Inches(3.15), Inches(2.5), Inches(0.5),
                      title, font_size=16, bold=True, color=MAGENTA)
        _add_text_box(slide, x + Inches(0.2), Inches(3.8), Inches(2.5), Inches(2),
                      desc, font_size=14, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 6 — Architecture diagram (text-based)
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "2. Architecture — Hybrid Strategy")

    # Pipeline flow
    steps = ["User Input", "LLM Agent", "Validator Gate", "FSM Memory", "JSON Output"]
    for idx, label in enumerate(steps):
        x = Inches(0.8 + idx * 2.5)
        fill = MAGENTA if idx in (1,) else SURFACE
        box = _add_shape_rect(slide, x, Inches(2.2), Inches(2), Inches(1), fill)
        _add_text_box(slide, x + Inches(0.1), Inches(2.35), Inches(1.8), Inches(0.7),
                      label, font_size=15, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
        if idx < len(steps) - 1:
            _add_text_box(slide, x + Inches(2), Inches(2.35), Inches(0.5), Inches(0.7),
                          "→", font_size=24, bold=True, color=MAGENTA, alignment=PP_ALIGN.CENTER)

    _add_text_box(slide, Inches(0.8), Inches(3.6), Inches(11), Inches(0.5),
                  "Every agent response is validated before the FSM advances — hallucinations never reach storage.",
                  font_size=16, color=LIGHT_GRAY)

    # Key tech stack
    _add_shape_rect(slide, Inches(0.8), Inches(4.5), Inches(11.5), Inches(2.2), SURFACE)
    _add_text_box(slide, Inches(1.1), Inches(4.6), Inches(5), Inches(0.5),
                  "Technology Stack", font_size=18, bold=True, color=MAGENTA)

    tech_items = [
        "Azure OpenAI (GPT-4o) — natural language understanding for 7 specialized agents",
        "15-State FSM — enforces valid transitions; no skipping steps",
        "Deterministic Validators — injection detection, SQL parsing, domain/subdomain checks",
        "Streamlit SPA — T-Mobile branded chat interface with buttons, forms & live JSON preview",
        "Backlog System — every table add/edit/delete is logged with linked summary + detail tables",
    ]
    _add_bullet_list(slide, Inches(1.1), Inches(5.2), Inches(10.5), Inches(2),
                     tech_items, font_size=14, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 7 — Security & Guardrails
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "2. Security & Guardrails")

    guards = [
        ("Prompt Injection Detection",
         "13 regex patterns block jailbreak attempts\n(\"ignore previous instructions\", DAN, etc.)"),
        ("Input Sanitisation",
         "HTML stripping, length limits (2000 chars),\nSQL mutation keyword blocking"),
        ("FSM Transition Enforcement",
         "Only legal state transitions are allowed.\nIllegal jumps are logged and blocked."),
        ("Isolated Agent Prompts",
         "Each of the 7 agents has its own system\nprompt — no cross-contamination."),
        ("Deterministic Gating",
         "LLM outputs always pass through validators\nbefore reaching the data model."),
        ("Backlog Audit Trail",
         "Every schema change is recorded with ID,\ntimestamp, action, before/after DDL."),
    ]
    for idx, (title, desc) in enumerate(guards):
        col = idx % 3
        row = idx // 3
        x = Inches(0.6 + col * 4.1)
        y = Inches(1.8 + row * 2.6)
        _add_shape_rect(slide, x, y, Inches(3.8), Inches(2.2), SURFACE)
        _add_text_box(slide, x + Inches(0.25), y + Inches(0.15), Inches(3.3), Inches(0.5),
                      title, font_size=15, bold=True, color=MAGENTA)
        _add_text_box(slide, x + Inches(0.25), y + Inches(0.75), Inches(3.3), Inches(1.3),
                      desc, font_size=13, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 8 — Demo intro
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "3. Live Demo")

    _add_text_box(slide, Inches(1), Inches(2.0), Inches(11), Inches(0.8),
                  "Scenario: Telecom Billing Bot", font_size=28, bold=True, color=WHITE)

    _add_text_box(slide, Inches(1), Inches(3.0), Inches(11), Inches(1.2),
                  "A telecom company needs a Text-to-SQL bot that lets support agents query billing data.\n"
                  "Today this takes a senior engineer 2+ hours to configure. With DMACO, a product manager\n"
                  "can do it in under 5 minutes — with full validation and an audit trail.",
                  font_size=18, color=LIGHT_GRAY)

    _add_shape_rect(slide, Inches(1), Inches(4.8), Inches(11), Inches(0.04), MAGENTA)

    _add_text_box(slide, Inches(1), Inches(5.1), Inches(11), Inches(0.5),
                  "What we'll walk through:", font_size=18, bold=True, color=MAGENTA)
    _add_bullet_list(slide, Inches(1), Inches(5.6), Inches(10), Inches(1.5), [
        "Select the Telecom domain → auto-loads domain prompt",
        "Pre-loaded billing tables (subscribers, invoices, payments) → add a custom table",
        "Choose Billing sub-domain → configure sub-domain prompt",
        "Add few-shot examples (natural language → SQL) → export final JSON",
    ], font_size=16, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 9 — Demo Step 1: Domain Selection
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "Demo Step 1 — Domain Selection")

    _add_shape_rect(slide, Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.5), SURFACE)
    _add_text_box(slide, Inches(1.1), Inches(1.95), Inches(5), Inches(0.5),
                  "What the user sees", font_size=16, bold=True, color=MAGENTA)
    _add_bullet_list(slide, Inches(1.1), Inches(2.6), Inches(5), Inches(3.5), [
        'Four domain buttons: Sales, Finance, Healthcare, Telecom',
        'User clicks "Telecom"',
        'Bot asks: "You chose Telecom. Confirm? (Yes / No)"',
        'User clicks "Yes" → moves to General Prompt',
    ], font_size=15, color=LIGHT_GRAY)

    _add_shape_rect(slide, Inches(7), Inches(1.8), Inches(5.5), Inches(4.5), SURFACE)
    _add_text_box(slide, Inches(7.3), Inches(1.95), Inches(5), Inches(0.5),
                  "What happens under the hood", font_size=16, bold=True, color=MAGENTA)
    _add_bullet_list(slide, Inches(7.3), Inches(2.6), Inches(5), Inches(3.5), [
        'FSM state: DOMAIN_SELECTION → DOMAIN_CONFIRMATION',
        'Domain agent (GPT-4o) interprets the choice',
        'Validator checks against allowed domains list',
        'On confirm: FSM advances to GENERAL_PROMPT_ENTRY',
        'Built-in prompt auto-loaded for Telecom',
    ], font_size=15, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 10 — Demo Step 2: Tables
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "Demo Step 2 — Table Configuration")

    _add_shape_rect(slide, Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.8), SURFACE)
    _add_text_box(slide, Inches(1.1), Inches(1.95), Inches(5), Inches(0.5),
                  "Pre-loaded tables (from mock data)", font_size=16, bold=True, color=MAGENTA)
    tables_text = (
        "subscribers\n"
        "  subscriber_id, phone_number, plan_name, status, ...\n\n"
        "invoices\n"
        "  invoice_id, subscriber_id, billing_period, amount_due, ...\n\n"
        "payments\n"
        "  payment_id, invoice_id, payment_date, amount_paid, ..."
    )
    _add_text_box(slide, Inches(1.1), Inches(2.6), Inches(5), Inches(3.5),
                  tables_text, font_size=14, color=LIGHT_GRAY)

    _add_shape_rect(slide, Inches(7), Inches(1.8), Inches(5.5), Inches(4.8), SURFACE)
    _add_text_box(slide, Inches(7.3), Inches(1.95), Inches(5), Inches(0.5),
                  "Adding a custom table", font_size=16, bold=True, color=MAGENTA)
    _add_text_box(slide, Inches(7.3), Inches(2.6), Inches(5), Inches(4), (
        "User clicks 'Add Table' and enters:\n\n"
        "CREATE TABLE usage_records (\n"
        "  record_id INT PRIMARY KEY,\n"
        "  subscriber_id INT NOT NULL,\n"
        "  data_mb DECIMAL(10,2),\n"
        "  voice_minutes INT,\n"
        "  sms_count INT,\n"
        "  record_date DATE NOT NULL\n"
        ");\n\n"
        "→ Validator parses DDL, extracts columns\n"
        "→ Backlog logs: \"Added table 'usage_records'\""
    ), font_size=14, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 11 — Demo Step 3: Sub-domain & Few-shots
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "Demo Step 3 — Sub-domain & Few-shot Examples")

    _add_shape_rect(slide, Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.8), SURFACE)
    _add_text_box(slide, Inches(1.1), Inches(1.95), Inches(5), Inches(0.5),
                  "Sub-domain: Billing", font_size=16, bold=True, color=MAGENTA)
    _add_bullet_list(slide, Inches(1.1), Inches(2.6), Inches(5), Inches(2), [
        'User picks "Billing" from 4 focus areas',
        'Auto-loaded prompt: "Focus on billing queries..."',
        'User clicks "Keep" to accept',
    ], font_size=15, color=LIGHT_GRAY)

    _add_text_box(slide, Inches(1.1), Inches(4.2), Inches(5), Inches(0.5),
                  "Pre-loaded few-shot examples", font_size=15, bold=True, color=MAGENTA)
    _add_text_box(slide, Inches(1.1), Inches(4.8), Inches(5), Inches(1.5), (
        'Q: "What is the total outstanding balance?"\n'
        'SQL: SELECT SUM(amount_due - amount_paid) ...\n\n'
        'Q: "Show overdue invoices for this month"\n'
        'SQL: SELECT * FROM invoices WHERE ...'
    ), font_size=13, color=LIGHT_GRAY)

    _add_shape_rect(slide, Inches(7), Inches(1.8), Inches(5.5), Inches(4.8), SURFACE)
    _add_text_box(slide, Inches(7.3), Inches(1.95), Inches(5), Inches(0.5),
                  "Adding a custom few-shot", font_size=16, bold=True, color=MAGENTA)
    _add_text_box(slide, Inches(7.3), Inches(2.6), Inches(5), Inches(4), (
        "User adds via structured form:\n\n"
        "Question:\n"
        '  "Show data usage per subscriber this month"\n\n'
        "SQL:\n"
        "  SELECT s.phone_number,\n"
        "         SUM(u.data_mb) AS total_data_mb\n"
        "  FROM usage_records u\n"
        "  JOIN subscribers s\n"
        "    ON u.subscriber_id = s.subscriber_id\n"
        "  WHERE u.record_date >= DATE_TRUNC(\n"
        "    'month', CURRENT_DATE)\n"
        "  GROUP BY s.phone_number;\n\n"
        "→ Validator cross-checks tables & columns\n"
        "→ References usage_records (custom table) ✓"
    ), font_size=13, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 12 — Demo Step 4: Export & Backlog
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "Demo Step 4 — Export & Backlog")

    _add_shape_rect(slide, Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.8), SURFACE)
    _add_text_box(slide, Inches(1.1), Inches(1.95), Inches(5), Inches(0.5),
                  "Final JSON Output", font_size=16, bold=True, color=MAGENTA)
    json_preview = (
        '{\n'
        '  "domain": "telecom",\n'
        '  "general_prompt": "You are an AI assistant...",\n'
        '  "tables": [\n'
        '    {"name": "subscribers", "description": "..."},\n'
        '    {"name": "invoices", "description": "..."},\n'
        '    {"name": "payments", "description": "..."},\n'
        '    {"name": "usage_records", "description": "..."}\n'
        '  ],\n'
        '  "sub_domain": "billing",\n'
        '  "sub_domain_prompt": "Focus on billing...",\n'
        '  "few_shots": [ ... 3 examples ... ]\n'
        '}'
    )
    _add_text_box(slide, Inches(1.1), Inches(2.6), Inches(5), Inches(4),
                  json_preview, font_size=13, color=LIGHT_GRAY)

    _add_shape_rect(slide, Inches(7), Inches(1.8), Inches(5.5), Inches(4.8), SURFACE)
    _add_text_box(slide, Inches(7.3), Inches(1.95), Inches(5), Inches(0.5),
                  "Backlog Audit Trail", font_size=16, bold=True, color=MAGENTA)
    _add_text_box(slide, Inches(7.3), Inches(2.6), Inches(5), Inches(1), (
        "Summary Table"
    ), font_size=14, bold=True, color=WHITE)
    _add_text_box(slide, Inches(7.3), Inches(3.1), Inches(5), Inches(1.2), (
        "ID            Date                        Summary\n"
        "a1b2c3d4  2026-04-08T14:30Z  Added table 'usage_records'"
    ), font_size=12, color=LIGHT_GRAY)

    _add_text_box(slide, Inches(7.3), Inches(4.3), Inches(5), Inches(0.5), (
        "Details Table"
    ), font_size=14, bold=True, color=WHITE)
    _add_text_box(slide, Inches(7.3), Inches(4.8), Inches(5), Inches(1.5), (
        "ID            Action   Table               DDL\n"
        "a1b2c3d4  add        usage_records   CREATE TABLE usage_records (...)\n\n"
        "→ Linked by ID — full traceability from summary to details"
    ), font_size=12, color=LIGHT_GRAY)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 13 — Key Takeaways
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _section_header(slide, "Key Takeaways")

    takeaways = [
        ("< 5 min", "Configuration time\n(vs. 2-4 hours manually)"),
        ("0 errors", "Validators catch issues\nbefore they reach storage"),
        ("Non-technical\nfriendly", "Product managers can\nconfigure bots directly"),
        ("Full audit trail", "Backlog tracks every\ntable change with ID + timestamp"),
    ]
    for idx, (big_num, desc) in enumerate(takeaways):
        x = Inches(0.6 + idx * 3.15)
        _add_shape_rect(slide, x, Inches(2.0), Inches(2.9), Inches(3.5), SURFACE)
        _add_text_box(slide, x + Inches(0.2), Inches(2.3), Inches(2.5), Inches(1),
                      big_num, font_size=28, bold=True, color=MAGENTA, alignment=PP_ALIGN.CENTER)
        _add_text_box(slide, x + Inches(0.2), Inches(3.6), Inches(2.5), Inches(1.5),
                      desc, font_size=15, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)

    # ══════════════════════════════════════════════════════════════════
    # SLIDE 14 — Thank You / Q&A
    # ══════════════════════════════════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_shape_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.12), MAGENTA)

    _add_text_box(slide, Inches(1), Inches(2.5), Inches(11), Inches(1),
                  "Thank You", font_size=48, bold=True, color=MAGENTA, alignment=PP_ALIGN.CENTER)
    _add_text_box(slide, Inches(1), Inches(3.8), Inches(11), Inches(0.8),
                  "Questions & Discussion", font_size=24, color=WHITE, alignment=PP_ALIGN.CENTER)
    _add_shape_rect(slide, Inches(5.5), Inches(4.8), Inches(2.3), Inches(0.04), MAGENTA)

    return prs


if __name__ == "__main__":
    prs = build_presentation()
    out = "DMACO_Presentation.pptx"
    prs.save(out)
    print(f"Saved → {out}")
